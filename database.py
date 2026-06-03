import io
import os
from contextlib import contextmanager
import psycopg2
import psycopg2.extras
import psycopg2.pool
import pandas as pd

DEFAULT_CATEGORIES = [
    ("Arriendo", "#FF6B6B"),
    ("Gastos comunes", "#FF8C42"),
    ("Comida", "#FFA07A"),
    ("Luz", "#FFD700"),
    ("Agua", "#5BC0EB"),
    ("Gas", "#9B5DE5"),
    ("Internet", "#00B4D8"),
    ("Teléfonos", "#06D6A0"),
    ("Transporte", "#4CC9F0"),
    ("Seguro de auto", "#F77F00"),
    ("Seguro de salud", "#D62246"),
    ("Higiene personal", "#4ECDC4"),
    ("Higiene hogar", "#45B7D1"),
    ("Bencina", "#FF6348"),
    ("Tag", "#2F86A6"),
    ("Salidas a comer", "#E84393"),
    ("Cultura", "#7B2D8B"),
    ("Bares", "#8B4513"),
    ("Café", "#C67C4E"),
    ("Suscripciones", "#6C757D"),
    ("Gym", "#F72585"),
    ("Miscellaneous", "#ADB5BD"),
    ("Ahorro para viajes", "#00D4FF"),
    ("Ahorro de emergencia", "#2DC653"),
    ("Ahorro para España", "#3A86FF"),
]

PERSON_NAMES = {"SG": "Santiago", "AZ": "Alex"}
PERSONS = ["SG", "AZ"]

MONTHS_ES = {
    1: "Enero", 2: "Febrero", 3: "Marzo", 4: "Abril",
    5: "Mayo", 6: "Junio", 7: "Julio", 8: "Agosto",
    9: "Septiembre", 10: "Octubre", 11: "Noviembre", 12: "Diciembre",
}

SPLIT_TYPES = {
    "personal":  "Personal",
    "shared":    "Compartido (50/50)",
    "for_other": "Para el otro",
    "custom":    "% Personalizado",
}


def _get_db_params() -> dict:
    try:
        import streamlit as st
        db = st.secrets["database"]
        return {
            "host":     db["host"],
            "port":     int(db["port"]),
            "dbname":   db["dbname"],
            "user":     db["user"],
            "password": db["password"],
            "sslmode":  "require",
        }
    except Exception:
        # Fallback: parse DATABASE_URL env var
        from urllib.parse import urlparse
        url = os.environ.get("DATABASE_URL", "")
        p = urlparse(url)
        return {
            "host":     p.hostname,
            "port":     p.port or 5432,
            "dbname":   p.path.lstrip("/"),
            "user":     p.username or "",
            "password": p.password or "",
            "sslmode":  "require",
        }


_pool = None

def _get_pool():
    global _pool
    if _pool is None or _pool.closed:
        _pool = psycopg2.pool.ThreadedConnectionPool(1, 5, **_get_db_params())
    return _pool


@contextmanager
def get_conn():
    pool = _get_pool()
    conn = pool.getconn()
    try:
        yield conn
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        try:
            pool.putconn(conn)
        except Exception:
            pass


def _run(conn, sql: str, params=()):
    """Execute a statement and return the cursor."""
    cur = conn.cursor()
    cur.execute(sql, params)
    return cur


def _df(conn, sql: str, params=()) -> pd.DataFrame:
    """Execute a SELECT and return a DataFrame."""
    cur = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cur.execute(sql, params)
    rows = cur.fetchall()
    return pd.DataFrame([dict(r) for r in rows])


# ─── Init ─────────────────────────────────────────────────────────────────────
_db_initialized = False

def init_db():
    global _db_initialized
    if _db_initialized:
        return
    with get_conn() as conn:
        _run(conn, """
            CREATE TABLE IF NOT EXISTS categories (
                id    SERIAL PRIMARY KEY,
                name  TEXT UNIQUE NOT NULL,
                color TEXT DEFAULT '#808080'
            )
        """)
        _run(conn, """
            CREATE TABLE IF NOT EXISTS budgets (
                id          SERIAL PRIMARY KEY,
                category_id INTEGER NOT NULL REFERENCES categories(id),
                person      TEXT NOT NULL,
                amount      REAL NOT NULL DEFAULT 0,
                month       INTEGER NOT NULL,
                year        INTEGER NOT NULL,
                UNIQUE(category_id, person, month, year)
            )
        """)
        _run(conn, """
            CREATE TABLE IF NOT EXISTS expenses (
                id            SERIAL PRIMARY KEY,
                description   TEXT NOT NULL,
                category_id   INTEGER NOT NULL REFERENCES categories(id),
                payer         TEXT NOT NULL,
                amount        REAL NOT NULL,
                split_type    TEXT NOT NULL DEFAULT 'personal',
                date          TEXT NOT NULL,
                budget_month  INTEGER,
                budget_year   INTEGER,
                is_reconciled INTEGER DEFAULT 0,
                notes         TEXT,
                created_at    TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)
        _run(conn, """
            CREATE TABLE IF NOT EXISTS settlements (
                id          SERIAL PRIMARY KEY,
                from_person TEXT NOT NULL,
                to_person   TEXT NOT NULL,
                amount      REAL NOT NULL,
                description TEXT,
                date        TEXT NOT NULL,
                created_at  TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)
        _run(conn, """
            CREATE TABLE IF NOT EXISTS manual_debts (
                id          SERIAL PRIMARY KEY,
                debtor      TEXT NOT NULL,
                creditor    TEXT NOT NULL,
                amount      REAL NOT NULL,
                description TEXT NOT NULL,
                date        TEXT NOT NULL,
                is_settled  INTEGER DEFAULT 0,
                created_at  TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)

        _run(conn, """
            CREATE TABLE IF NOT EXISTS savings (
                id          SERIAL PRIMARY KEY,
                person      TEXT NOT NULL,
                amount      REAL NOT NULL,
                entry_type  TEXT NOT NULL DEFAULT 'deposit',
                description TEXT NOT NULL,
                date        TEXT NOT NULL,
                created_at  TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)

        _run(conn, """
            CREATE TABLE IF NOT EXISTS income_monthly (
                id      SERIAL PRIMARY KEY,
                person  TEXT NOT NULL,
                amount  REAL NOT NULL DEFAULT 0,
                month   INTEGER NOT NULL,
                year    INTEGER NOT NULL,
                UNIQUE(person, month, year)
            )
        """)
        _run(conn, """
            CREATE TABLE IF NOT EXISTS income_extra (
                id          SERIAL PRIMARY KEY,
                person      TEXT NOT NULL,
                description TEXT NOT NULL,
                amount      REAL NOT NULL,
                date        TEXT NOT NULL,
                notes       TEXT,
                created_at  TEXT DEFAULT CURRENT_TIMESTAMP
            )
        """)

        # Safe column migrations for existing databases
        for col, definition in [
            ("budget_month",  "INTEGER"),
            ("budget_year",   "INTEGER"),
            ("is_reconciled", "INTEGER DEFAULT 0"),
            ("split_pct",     "REAL"),
        ]:
            _run(conn, f"ALTER TABLE expenses ADD COLUMN IF NOT EXISTS {col} {definition}")
        _run(conn, "ALTER TABLE settlements ADD COLUMN IF NOT EXISTS debt_type TEXT DEFAULT 'period'")
        _run(conn, "ALTER TABLE savings ADD COLUMN IF NOT EXISTS expense_id INTEGER")

        for old, new in [
            ("Emergency Savings", "Ahorro de emergencia"),
            ("Spain Move Fund",   "Ahorro para España"),
            ("Viajes",            "Ahorro para viajes"),
        ]:
            _run(conn, "UPDATE categories SET name=%s WHERE name=%s", (new, old))

        for name, color in DEFAULT_CATEGORIES:
            _run(conn,
                 "INSERT INTO categories (name, color) VALUES (%s, %s) ON CONFLICT (name) DO NOTHING",
                 (name, color))
    _db_initialized = True


# ─── Categories ───────────────────────────────────────────────────────────────
def get_categories() -> pd.DataFrame:
    with get_conn() as conn:
        return _df(conn, "SELECT * FROM categories ORDER BY name")


def add_category(name: str, color: str = "#808080") -> tuple:
    name = name.strip()
    if not name:
        return False, "El nombre no puede estar vacío."
    try:
        with get_conn() as conn:
            _run(conn, "INSERT INTO categories (name, color) VALUES (%s, %s)", (name, color))
        return True, f'Categoría "{name}" creada.'
    except psycopg2.IntegrityError:
        return False, f'La categoría "{name}" ya existe.'


def delete_category(cat_id: int) -> tuple:
    with get_conn() as conn:
        cur = _run(conn, "SELECT COUNT(*) FROM expenses WHERE category_id=%s", (cat_id,))
        n = cur.fetchone()[0]
        if n > 0:
            return False, f"No se puede eliminar: tiene {n} gasto(s) asociado(s)."
        _run(conn, "DELETE FROM budgets    WHERE category_id=%s", (cat_id,))
        _run(conn, "DELETE FROM categories WHERE id=%s",          (cat_id,))
    return True, "Categoría eliminada."


def update_category(cat_id: int, name: str, color: str) -> tuple:
    name = name.strip()
    if not name:
        return False, "El nombre no puede estar vacío."
    try:
        with get_conn() as conn:
            _run(conn, "UPDATE categories SET name=%s, color=%s WHERE id=%s", (name, color, cat_id))
        return True, f'Categoría actualizada a "{name}".'
    except psycopg2.IntegrityError:
        return False, f'Ya existe una categoría llamada "{name}".'


# ─── Budgets ──────────────────────────────────────────────────────────────────
def set_budget(category_id: int, person: str, amount: float, month: int, year: int):
    with get_conn() as conn:
        _run(conn, """
            INSERT INTO budgets (category_id, person, amount, month, year)
            VALUES (%s, %s, %s, %s, %s)
            ON CONFLICT(category_id, person, month, year)
            DO UPDATE SET amount = EXCLUDED.amount
        """, (category_id, person, amount, month, year))


def get_budgets(month: int, year: int) -> pd.DataFrame:
    cats_df = get_categories()
    with get_conn() as conn:
        budgets_raw = _df(conn,
            "SELECT category_id, person, amount FROM budgets WHERE month=%s AND year=%s",
            (month, year))
    rows = []
    for _, cat in cats_df.iterrows():
        row = {"category_id": int(cat["id"]), "category_name": cat["name"], "color": cat["color"]}
        for person in PERSONS:
            if not budgets_raw.empty:
                match = budgets_raw[
                    (budgets_raw["category_id"] == cat["id"]) & (budgets_raw["person"] == person)
                ]
                row[f"budget_{person}"] = float(match["amount"].iloc[0]) if not match.empty else 0.0
            else:
                row[f"budget_{person}"] = 0.0
        rows.append(row)
    return pd.DataFrame(rows)


def copy_budgets_from_month(from_month: int, from_year: int, to_month: int, to_year: int):
    with get_conn() as conn:
        cur = _run(conn,
            "SELECT category_id, person, amount FROM budgets WHERE month=%s AND year=%s",
            (from_month, from_year))
        rows = cur.fetchall()
        for r in rows:
            _run(conn, """
                INSERT INTO budgets (category_id, person, amount, month, year)
                VALUES (%s, %s, %s, %s, %s)
                ON CONFLICT(category_id, person, month, year)
                DO UPDATE SET amount = EXCLUDED.amount
            """, (r[0], r[1], r[2], to_month, to_year))


# ─── Expenses ─────────────────────────────────────────────────────────────────
def add_expense(
    description: str,
    category_id: int,
    payer: str,
    amount: float,
    split_type: str,
    expense_date: str,
    notes: str = None,
    budget_month: int = None,
    budget_year: int = None,
    split_pct: float = None,
):
    with get_conn() as conn:
        cur = _run(conn, """
            INSERT INTO expenses
              (description, category_id, payer, amount, split_type, date, notes, budget_month, budget_year, split_pct)
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            RETURNING id
        """, (description, category_id, payer, amount, split_type, expense_date, notes, budget_month, budget_year, split_pct))
        return int(cur.fetchone()[0])


def get_expenses(month: int = None, year: int = None) -> pd.DataFrame:
    with get_conn() as conn:
        query = """
            SELECT e.id, e.description, e.payer, e.amount, e.split_type,
                   e.date, e.budget_month, e.budget_year, e.notes, e.created_at,
                   e.is_reconciled, e.split_pct,
                   c.id as category_id, c.name as category_name, c.color
            FROM expenses e
            JOIN categories c ON e.category_id = c.id
        """
        params: list = []
        if month is not None and year is not None:
            query += """
                WHERE (
                    (e.budget_month IS NOT NULL AND e.budget_month = %s AND e.budget_year = %s)
                    OR
                    (e.budget_month IS NULL
                     AND EXTRACT(MONTH FROM e.date::date) = %s
                     AND EXTRACT(YEAR  FROM e.date::date) = %s)
                )
            """
            params = [month, year, month, year]
        query += " ORDER BY e.date DESC, e.created_at DESC"
        return _df(conn, query, params)


def update_expense(
    expense_id: int,
    description: str,
    category_id: int,
    payer: str,
    amount: float,
    split_type: str,
    expense_date: str,
    notes: str = None,
    budget_month: int = None,
    budget_year: int = None,
    split_pct: float = None,
):
    with get_conn() as conn:
        _run(conn, """
            UPDATE expenses SET
                description=%s, category_id=%s, payer=%s, amount=%s,
                split_type=%s, date=%s, notes=%s, budget_month=%s,
                budget_year=%s, split_pct=%s
            WHERE id=%s
        """, (description, category_id, payer, amount, split_type,
              expense_date, notes, budget_month, budget_year, split_pct, expense_id))


def delete_savings_by_expense(expense_id: int):
    with get_conn() as conn:
        _run(conn, "DELETE FROM savings WHERE expense_id=%s", (expense_id,))


def delete_expense(expense_id: int):
    with get_conn() as conn:
        _run(conn, "DELETE FROM savings  WHERE expense_id=%s", (expense_id,))
        _run(conn, "DELETE FROM expenses WHERE id=%s",         (expense_id,))


def reconcile_expense(expense_id: int, reconciled: bool = True):
    with get_conn() as conn:
        _run(conn, "UPDATE expenses SET is_reconciled=%s WHERE id=%s",
             (1 if reconciled else 0, expense_id))


# ─── Settlements ──────────────────────────────────────────────────────────────
def add_settlement(from_person: str, to_person: str, amount: float, description: str,
                   settlement_date: str, debt_type: str = "period"):
    with get_conn() as conn:
        _run(conn,
             "INSERT INTO settlements (from_person, to_person, amount, description, date, debt_type) VALUES (%s,%s,%s,%s,%s,%s)",
             (from_person, to_person, amount, description, settlement_date, debt_type))


def get_settlements(month: int = None, year: int = None) -> pd.DataFrame:
    with get_conn() as conn:
        query = "SELECT * FROM settlements"
        params: list = []
        if month is not None and year is not None:
            query += (" WHERE EXTRACT(MONTH FROM date::date)=%s"
                      "   AND EXTRACT(YEAR  FROM date::date)=%s")
            params = [month, year]
        query += " ORDER BY date DESC"
        return _df(conn, query, params)


def delete_settlement(settlement_id: int):
    with get_conn() as conn:
        _run(conn, "DELETE FROM settlements WHERE id=%s", (settlement_id,))


# ─── Manual Debts ─────────────────────────────────────────────────────────────
def add_manual_debt(debtor: str, creditor: str, amount: float, description: str, debt_date: str):
    with get_conn() as conn:
        _run(conn,
             "INSERT INTO manual_debts (debtor, creditor, amount, description, date) VALUES (%s,%s,%s,%s,%s)",
             (debtor, creditor, amount, description, debt_date))


def get_manual_debts(only_pending: bool = True) -> pd.DataFrame:
    with get_conn() as conn:
        query = "SELECT * FROM manual_debts"
        if only_pending:
            query += " WHERE is_settled = 0"
        query += " ORDER BY date DESC"
        return _df(conn, query)


def settle_manual_debt(debt_id: int):
    with get_conn() as conn:
        _run(conn, "UPDATE manual_debts SET is_settled = 1 WHERE id = %s", (debt_id,))


def delete_manual_debt(debt_id: int):
    with get_conn() as conn:
        _run(conn, "DELETE FROM manual_debts WHERE id=%s", (debt_id,))


# ─── Analytics ────────────────────────────────────────────────────────────────
def _other_pct(exp) -> float:
    """Fraction of the expense that the other person owes/pays."""
    st = exp["split_type"]
    if st == "shared":    return 0.5
    if st == "for_other": return 1.0
    if st == "custom":
        v = exp.get("split_pct")
        return float(v) / 100.0 if v is not None and not pd.isna(v) else 0.5
    return 0.0  # personal


def calculate_spending_by_person_category(month: int, year: int,
                                           expenses: pd.DataFrame = None) -> pd.DataFrame:
    if expenses is None:
        expenses = get_expenses(month=month, year=year)
    if expenses.empty:
        return pd.DataFrame(columns=["category_id", "category_name", "color", "person", "spent"])

    rows = []
    for _, exp in expenses.iterrows():
        cat   = {"category_id": exp["category_id"], "category_name": exp["category_name"], "color": exp["color"]}
        amt   = float(exp["amount"])
        other = "AZ" if exp["payer"] == "SG" else "SG"
        opct  = _other_pct(exp)
        if exp["split_type"] == "personal":
            rows.append({**cat, "person": exp["payer"], "spent": amt})
        else:
            rows.append({**cat, "person": exp["payer"], "spent": amt * (1 - opct)})
            rows.append({**cat, "person": other,        "spent": amt * opct})

    df = pd.DataFrame(rows)
    return df.groupby(["category_id", "category_name", "color", "person"])["spent"].sum().reset_index()


def calculate_period_balance(month: int, year: int,
                              expenses: pd.DataFrame = None,
                              settlements: pd.DataFrame = None) -> float:
    """Balance from UNRECONCILED shared expenses this month, minus period settlements.
    Positive = AZ owes SG. Negative = SG owes AZ."""
    if expenses    is None: expenses    = get_expenses(month=month, year=year)
    if settlements is None: settlements = get_settlements(month=month, year=year)
    balance = 0.0

    if not expenses.empty:
        for _, row in expenses[expenses["is_reconciled"].fillna(0) != 1].iterrows():
            if row["split_type"] in ("shared", "for_other", "custom"):
                amt = float(row["amount"]) * _other_pct(row)
                balance += amt if row["payer"] == "SG" else -amt

    if not settlements.empty:
        for _, row in settlements[settlements.get("debt_type", "period") == "period"].iterrows():
            amt = float(row["amount"])
            if row["from_person"] == "AZ" and row["to_person"] == "SG":
                balance -= amt
            elif row["from_person"] == "SG" and row["to_person"] == "AZ":
                balance += amt

    return balance


def calculate_accumulated_balance(reconciled_expenses: pd.DataFrame = None,
                                   manual: pd.DataFrame = None,
                                   accumulated_settlements: pd.DataFrame = None) -> float:
    """Balance from RECONCILED shared expenses (all months) + manual debts + accumulated payments.
    Positive = AZ owes SG. Negative = SG owes AZ."""
    if reconciled_expenses is None:
        all_exp = get_expenses()
        reconciled_expenses = (
            all_exp[all_exp["is_reconciled"].fillna(0) == 1]
            if not all_exp.empty else pd.DataFrame()
        )
    if manual is None:
        manual = get_manual_debts(only_pending=True)
    if accumulated_settlements is None:
        all_sett = get_settlements()
        accumulated_settlements = (
            all_sett[all_sett["debt_type"] == "accumulated"]
            if not all_sett.empty and "debt_type" in all_sett.columns else pd.DataFrame()
        )
    balance = 0.0

    if not reconciled_expenses.empty:
        for _, row in reconciled_expenses.iterrows():
            if row["split_type"] in ("shared", "for_other", "custom"):
                amt = float(row["amount"]) * _other_pct(row)
                balance += amt if row["payer"] == "SG" else -amt

    if not manual.empty:
        for _, row in manual.iterrows():
            amt = float(row["amount"])
            if row["debtor"] == "AZ" and row["creditor"] == "SG":
                balance += amt
            elif row["debtor"] == "SG" and row["creditor"] == "AZ":
                balance -= amt

    if not accumulated_settlements.empty:
        for _, row in accumulated_settlements.iterrows():
            amt = float(row["amount"])
            if row["from_person"] == "AZ" and row["to_person"] == "SG":
                balance -= amt
            elif row["from_person"] == "SG" and row["to_person"] == "AZ":
                balance += amt

    return balance


def get_daily_spending(month: int, year: int, expenses: pd.DataFrame = None) -> pd.DataFrame:
    if expenses is None:
        expenses = get_expenses(month=month, year=year)
    if expenses.empty:
        return pd.DataFrame(columns=["date", "person", "amount"])

    rows = []
    for _, exp in expenses.iterrows():
        amt   = float(exp["amount"])
        other = "AZ" if exp["payer"] == "SG" else "SG"
        opct  = _other_pct(exp)
        if exp["split_type"] == "personal":
            rows.append({"date": exp["date"], "person": exp["payer"], "amount": amt})
        else:
            rows.append({"date": exp["date"], "person": exp["payer"], "amount": amt * (1 - opct)})
            rows.append({"date": exp["date"], "person": other,        "amount": amt * opct})

    df = pd.DataFrame(rows)
    result = df.groupby(["date", "person"])["amount"].sum().reset_index()
    result["date"] = pd.to_datetime(result["date"])
    return result.sort_values("date")


# ─── Savings ──────────────────────────────────────────────────────────────────
def add_savings_entry(person: str, amount: float, entry_type: str, description: str,
                      entry_date: str, expense_id: int = None):
    """entry_type: 'deposit' | 'previous' | 'withdrawal' | 'return' | 'loss'"""
    with get_conn() as conn:
        _run(conn,
             "INSERT INTO savings (person, amount, entry_type, description, date, expense_id) VALUES (%s,%s,%s,%s,%s,%s)",
             (person, amount, entry_type, description, entry_date, expense_id))


def get_savings(person: str = None) -> pd.DataFrame:
    with get_conn() as conn:
        if person:
            return _df(conn, "SELECT * FROM savings WHERE person=%s ORDER BY date DESC, created_at DESC", (person,))
        return _df(conn, "SELECT * FROM savings ORDER BY date DESC, created_at DESC")


def get_savings_balance(person: str) -> float:
    with get_conn() as conn:
        cur = _run(conn,
                   "SELECT COALESCE(SUM(CASE WHEN entry_type IN ('withdrawal','loss') THEN -amount ELSE amount END),0) FROM savings WHERE person=%s",
                   (person,))
        return float(cur.fetchone()[0])


def delete_savings_entry(entry_id: int):
    with get_conn() as conn:
        _run(conn, "DELETE FROM savings WHERE id=%s", (entry_id,))


# ─── Income ───────────────────────────────────────────────────────────────────
def set_monthly_income(person: str, amount: float, month: int, year: int):
    with get_conn() as conn:
        _run(conn, """
            INSERT INTO income_monthly (person, amount, month, year)
            VALUES (%s, %s, %s, %s)
            ON CONFLICT(person, month, year)
            DO UPDATE SET amount = EXCLUDED.amount
        """, (person, amount, month, year))


def get_monthly_income(month: int, year: int) -> dict:
    with get_conn() as conn:
        rows = _df(conn,
            "SELECT person, amount FROM income_monthly WHERE month=%s AND year=%s",
            (month, year))
    result = {p: 0.0 for p in PERSONS}
    if not rows.empty:
        for _, r in rows.iterrows():
            result[r["person"]] = float(r["amount"])
    return result


def add_income_entry(person: str, description: str, amount: float,
                     entry_date: str, notes: str = None) -> int:
    with get_conn() as conn:
        cur = _run(conn,
            "INSERT INTO income_extra (person, description, amount, date, notes) "
            "VALUES (%s,%s,%s,%s,%s) RETURNING id",
            (person, description, amount, entry_date, notes))
        return int(cur.fetchone()[0])


def get_income_entries(month: int = None, year: int = None) -> pd.DataFrame:
    with get_conn() as conn:
        query = "SELECT * FROM income_extra"
        params: list = []
        if month is not None and year is not None:
            query += (" WHERE EXTRACT(MONTH FROM date::date)=%s"
                      "   AND EXTRACT(YEAR  FROM date::date)=%s")
            params = [month, year]
        query += " ORDER BY date DESC, created_at DESC"
        return _df(conn, query, params)


def update_income_entry(entry_id: int, description: str, amount: float,
                        entry_date: str, notes: str = None):
    with get_conn() as conn:
        _run(conn,
            "UPDATE income_extra SET description=%s, amount=%s, date=%s, notes=%s WHERE id=%s",
            (description, amount, entry_date, notes, entry_id))


def delete_income_entry(entry_id: int):
    with get_conn() as conn:
        _run(conn, "DELETE FROM income_extra WHERE id=%s", (entry_id,))


# ─── Excel Export ─────────────────────────────────────────────────────────────
def build_excel_export(month: int, year: int) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    month_name = MONTHS_ES.get(month, str(month))

    HEADER_FILL = PatternFill("solid", fgColor="667EEA")
    HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
    TOTAL_FILL  = PatternFill("solid", fgColor="EEF2FF")
    TOTAL_FONT  = Font(bold=True, size=11)
    CENTER      = Alignment(horizontal="center", vertical="center")

    def style_header(ws):
        for cell in ws[1]:
            cell.fill      = HEADER_FILL
            cell.font      = HEADER_FONT
            cell.alignment = CENTER

    def autofit(ws):
        for col_cells in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col_cells), default=8)
            ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max_len + 4, 40)

    # ── Sheet 1: Gastos ───────────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = f"Gastos {month_name}"

    expenses = get_expenses(month=month, year=year)
    headers1 = ["Fecha", "Descripción", "Categoría", "Pagó", "Monto ($)", "Tipo", "Mes Presupuesto", "Notas"]
    ws1.append(headers1)
    style_header(ws1)

    if not expenses.empty:
        for _, r in expenses.iterrows():
            bm = ""
            if r.get("budget_month") and not (isinstance(r["budget_month"], float) and pd.isna(r["budget_month"])):
                bm = f"{MONTHS_ES.get(int(r['budget_month']), '?')} {int(r['budget_year'])}"
            ws1.append([
                r["date"], r["description"], r["category_name"],
                f"{r['payer']} · {PERSON_NAMES[r['payer']]}",
                round(float(r["amount"]), 2),
                SPLIT_TYPES.get(r["split_type"], r["split_type"]),
                bm, r["notes"] or "",
            ])
        last = ws1.max_row
        ws1.cell(last + 1, 1, "TOTAL").font = TOTAL_FONT
        ws1.cell(last + 1, 1).fill = TOTAL_FILL
        c = ws1.cell(last + 1, 5, round(expenses["amount"].sum(), 2))
        c.font = TOTAL_FONT
        c.fill = TOTAL_FILL

    autofit(ws1)
    ws1.freeze_panes = "A2"

    # ── Sheet 2: Presupuesto ──────────────────────────────────────────────────
    ws2 = wb.create_sheet(f"Presupuesto {month_name}")
    headers2 = [
        "Categoría",
        "Presupuesto SG ($)", "Gastado SG ($)", "Restante SG ($)",
        "Presupuesto AZ ($)", "Gastado AZ ($)", "Restante AZ ($)",
        "Total Presupuesto ($)", "Total Gastado ($)",
    ]
    ws2.append(headers2)
    style_header(ws2)

    budgets  = get_budgets(month=month, year=year)
    spending = calculate_spending_by_person_category(month=month, year=year)
    totals   = {k: 0.0 for k in ["bSG","sSG","rSG","bAZ","sAZ","rAZ","bT","sT"]}

    for _, brow in budgets.iterrows():
        row_data = [brow["category_name"]]
        for person, keys in [("SG", ("bSG","sSG","rSG")), ("AZ", ("bAZ","sAZ","rAZ"))]:
            budget    = brow[f"budget_{person}"]
            sp        = spending[(spending["person"] == person) & (spending["category_name"] == brow["category_name"])]
            spent     = sp["spent"].sum() if not sp.empty else 0.0
            remaining = budget - spent
            row_data += [round(budget, 2), round(spent, 2), round(remaining, 2)]
            totals[keys[0]] += budget
            totals[keys[1]] += spent
            totals[keys[2]] += remaining
        total_b = brow["budget_SG"] + brow["budget_AZ"]
        sp_all  = spending[spending["category_name"] == brow["category_name"]]
        total_s = sp_all["spent"].sum() if not sp_all.empty else 0.0
        row_data += [round(total_b, 2), round(total_s, 2)]
        totals["bT"] += total_b
        totals["sT"] += total_s
        ws2.append(row_data)

    tr = ws2.max_row + 1
    ws2.cell(tr, 1, "TOTAL").font = TOTAL_FONT
    ws2.cell(tr, 1).fill = TOTAL_FILL
    for i, val in enumerate([totals["bSG"], totals["sSG"], totals["rSG"],
                              totals["bAZ"], totals["sAZ"], totals["rAZ"],
                              totals["bT"],  totals["sT"]], start=2):
        c = ws2.cell(tr, i, round(val, 2))
        c.font = TOTAL_FONT
        c.fill = TOTAL_FILL

    autofit(ws2)
    ws2.freeze_panes = "A2"

    # ── Sheet 3: Deudas ───────────────────────────────────────────────────────
    ws3 = wb.create_sheet("Deudas")
    balance = calculate_period_balance(month, year) + calculate_accumulated_balance()
    ws3.append(["Resumen de Deudas", f"{month_name} {year}"])
    ws3["A1"].font = Font(bold=True, size=13)
    ws3.append([])
    abs_bal = abs(balance)
    if balance > 0.01:
        ws3.append(["Alex debe a Santiago:", f"${abs_bal:,.2f}"])
    elif balance < -0.01:
        ws3.append(["Santiago debe a Alex:", f"${abs_bal:,.2f}"])
    else:
        ws3.append(["Balance:", "Sin deudas"])
    ws3.append([])

    ws3.append(["Gastos Compartidos y Para el Otro"])
    ws3[f"A{ws3.max_row}"].font = Font(bold=True)
    headers3 = ["Fecha", "Descripción", "Categoría", "Pagó", "Monto ($)", "Tipo", "Deuda generada"]
    ws3.append(headers3)
    style_header(ws3)

    if not expenses.empty:
        shared = expenses[expenses["split_type"].isin(["shared", "for_other"])]
        for _, r in shared.iterrows():
            other    = "AZ" if r["payer"] == "SG" else "SG"
            debt_amt = r["amount"] / 2 if r["split_type"] == "shared" else r["amount"]
            ws3.append([
                r["date"], r["description"], r["category_name"], r["payer"],
                round(float(r["amount"]), 2),
                SPLIT_TYPES.get(r["split_type"], r["split_type"]),
                f"{other} debe ${debt_amt:,.2f}",
            ])
    ws3.append([])

    manual = get_manual_debts(only_pending=True)
    ws3.append(["Deudas Manuales Pendientes"])
    ws3[f"A{ws3.max_row}"].font = Font(bold=True)
    ws3.append(["Fecha", "Descripción", "Quien debe", "A quien", "Monto ($)"])
    if not manual.empty:
        for _, r in manual.iterrows():
            ws3.append([r["date"], r["description"],
                        f"{r['debtor']} · {PERSON_NAMES[r['debtor']]}",
                        f"{r['creditor']} · {PERSON_NAMES[r['creditor']]}",
                        round(float(r["amount"]), 2)])
    ws3.append([])

    settlements = get_settlements(month=month, year=year)
    ws3.append(["Pagos Registrados"])
    ws3[f"A{ws3.max_row}"].font = Font(bold=True)
    ws3.append(["Fecha", "De", "A", "Monto ($)", "Descripción"])
    if not settlements.empty:
        for _, r in settlements.iterrows():
            ws3.append([r["date"], r["from_person"], r["to_person"],
                        round(float(r["amount"]), 2), r["description"] or ""])

    autofit(ws3)

    # ── Sheet 4: Ahorros ──────────────────────────────────────────────────────
    ws4 = wb.create_sheet("Ahorros")
    ENTRY_LABELS_XL = {
        "deposit": "Depósito", "previous": "Ahorro previo",
        "withdrawal": "Retiro", "return": "Rentabilidad +", "loss": "Rentabilidad −",
    }
    for person in PERSONS:
        ws4.append([f"Ahorros — {PERSON_NAMES[person]}"])
        ws4[f"A{ws4.max_row}"].font = Font(bold=True, size=12)
        ws4.append(["Fecha", "Descripción", "Tipo", "Monto ($)"])
        style_header(ws4)
        sav = get_savings(person)
        if not sav.empty:
            for _, r in sav.iterrows():
                sign = -1 if r["entry_type"] in ("withdrawal", "loss") else 1
                ws4.append([r["date"], r["description"],
                            ENTRY_LABELS_XL.get(r["entry_type"], r["entry_type"]),
                            round(float(r["amount"]) * sign, 2)])
        balance_sv = get_savings_balance(person)
        tr = ws4.max_row + 1
        ws4.cell(tr, 1, "SALDO TOTAL").font = TOTAL_FONT
        ws4.cell(tr, 1).fill = TOTAL_FILL
        c = ws4.cell(tr, 4, round(balance_sv, 2))
        c.font = TOTAL_FONT
        c.fill = TOTAL_FILL
        ws4.append([])

    autofit(ws4)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


# ─── Santiago personal finances ───────────────────────────────────────────────

_SG_DEFAULT_ACCOUNTS = [
    ("Banco de Chile", "bank"),
    ("Santander",      "bank"),
    ("Tarjeta Santander", "credit"),
    ("CMR",            "credit"),
]

def init_sg_tables():
    with get_conn() as conn:
        _run(conn, """
            CREATE TABLE IF NOT EXISTS sg_accounts (
                id           SERIAL PRIMARY KEY,
                account_name VARCHAR(100) UNIQUE NOT NULL,
                account_type VARCHAR(20)  NOT NULL,
                balance      DECIMAL(12,2) NOT NULL DEFAULT 0,
                updated_at   TIMESTAMP DEFAULT NOW()
            )
        """)
        _run(conn, """
            CREATE TABLE IF NOT EXISTS sg_personal_debts (
                id          SERIAL PRIMARY KEY,
                person_name VARCHAR(100) NOT NULL,
                amount      DECIMAL(12,2) NOT NULL,
                description TEXT,
                date        DATE NOT NULL,
                status      VARCHAR(20) DEFAULT 'pending',
                paid_date   DATE,
                created_at  TIMESTAMP DEFAULT NOW()
            )
        """)
        for name, atype in _SG_DEFAULT_ACCOUNTS:
            _run(conn, """
                INSERT INTO sg_accounts (account_name, account_type, balance)
                VALUES (%s, %s, 0) ON CONFLICT (account_name) DO NOTHING
            """, (name, atype))

def get_sg_accounts() -> pd.DataFrame:
    with get_conn() as conn:
        return _df(conn, "SELECT * FROM sg_accounts ORDER BY account_type DESC, account_name")

def update_sg_account_balance(account_id: int, balance: float):
    with get_conn() as conn:
        _run(conn, "UPDATE sg_accounts SET balance=%s, updated_at=NOW() WHERE id=%s",
             (balance, account_id))

def get_sg_personal_debts(only_pending: bool = True) -> pd.DataFrame:
    with get_conn() as conn:
        q = ("SELECT * FROM sg_personal_debts WHERE status='pending' ORDER BY date DESC"
             if only_pending else
             "SELECT * FROM sg_personal_debts ORDER BY date DESC")
        return _df(conn, q)

def add_sg_personal_debt(person_name: str, amount: float, description: str, debt_date: str):
    with get_conn() as conn:
        _run(conn, """
            INSERT INTO sg_personal_debts (person_name, amount, description, date)
            VALUES (%s, %s, %s, %s)
        """, (person_name, amount, description, debt_date))

def settle_sg_personal_debt(debt_id: int, paid_date: str):
    with get_conn() as conn:
        _run(conn, "UPDATE sg_personal_debts SET status='paid', paid_date=%s WHERE id=%s",
             (paid_date, debt_id))

def update_sg_personal_debt_amount(debt_id: int, amount: float):
    with get_conn() as conn:
        _run(conn, "UPDATE sg_personal_debts SET amount=%s WHERE id=%s", (amount, debt_id))

def delete_sg_personal_debt(debt_id: int):
    with get_conn() as conn:
        _run(conn, "DELETE FROM sg_personal_debts WHERE id=%s", (debt_id,))
